# SPDX-FileCopyrightText: 2024-2026 Dan J. Bower <dbower@eaps.ethz.ch>
# SPDX-FileCopyrightText: 2024-2026 Paolo A. Sossi
#
# SPDX-License-Identifier: Apache-2.0

"""Tests for the public output API (:mod:`atmodeller.output`) and its base infrastructure
(:mod:`atmodeller.output_base`).

Three models are solved once each: a gas-only C-H-O model batched over two H budgets, an H-only
model (no heavy elements) and a C-O model (no H). Each output format is computed once per model.
"""

import logging
import pickle
from pathlib import Path
from typing import Any, Literal, Optional

import equinox as eqx
import jax.numpy as jnp
import numpy as np
import pandas as pd
import pytest
from jaxtyping import ArrayLike
from openpyxl import load_workbook

from atmodeller import __version__, debug_logger
from atmodeller.classes import EquilibriumModel
from atmodeller.containers import ChemicalSpecies
from atmodeller.interfaces import ActivityConstraintProtocol
from atmodeller.output import Output
from atmodeller.output_base import flatten_dictionary, log_epsilon, recursively_merge_dictionaries
from atmodeller.parameters import Parameters
from atmodeller.sci_utils import earth, solar
from atmodeller.state import BaseThermodynamicState, Planet
from atmodeller.thermodata import IronWustiteBuffer

logger: logging.Logger = debug_logger()
logger.setLevel(logging.DEBUG)

OutputFormat = Literal["natural", "named_arrays", "elements_species"]
OUTPUT_FORMATS: tuple[OutputFormat, ...] = ("natural", "named_arrays", "elements_species")
"""All output formats"""

# Gas species
H2O_g: ChemicalSpecies = ChemicalSpecies.create_gas("H2O")
H2_g: ChemicalSpecies = ChemicalSpecies.create_gas("H2")
H_g: ChemicalSpecies = ChemicalSpecies.create_gas("H")
O2_g: ChemicalSpecies = ChemicalSpecies.create_gas("O2")
CO_g: ChemicalSpecies = ChemicalSpecies.create_gas("CO")
CO2_g: ChemicalSpecies = ChemicalSpecies.create_gas("CO2")
CH4_g: ChemicalSpecies = ChemicalSpecies.create_gas("CH4")


def _solve(
    gas_species: tuple[ChemicalSpecies, ...],
    mass_constraints: dict[str, ArrayLike],
    activity_constraints: Optional[dict[str, ActivityConstraintProtocol]] = None,
) -> Output:
    """Builds and solves a model"""
    planet: BaseThermodynamicState = Planet.from_species(gas_species)
    parameters: Parameters = Parameters(
        planet, activity_constraints=activity_constraints, mass_constraints=mass_constraints
    )

    return EquilibriumModel(parameters).solve_with_default()


def _to_dicts(output: Output) -> dict[str, dict[str, Any]]:
    """Output in each format"""
    return {fmt: output.to_dict(output_format=fmt, to_numpy=True) for fmt in OUTPUT_FORMATS}


@pytest.fixture(scope="module")
def cho_output() -> Output:
    """Gas-only C-H-O model for two H budgets (C = H by mass)

    The first budget (one ocean) is compared with ``main``: with no dissolution or condensates the
    gas-phase inventory equals the total inventory. The batch of two exercises broadcasting of
    unbatched leaves, such as the temperature, in
    :func:`~atmodeller.output_base.expand_jax_arrays_in_pytree_to_batch`.
    """
    h_kg: ArrayLike = earth.oceans_to_hydrogen_mass(np.array([1.0, 2.0]))

    return _solve(
        (H2_g, H2O_g, CO_g, CO2_g, CH4_g, O2_g),
        {"H": h_kg, "C": h_kg},
        activity_constraints={"O2_g": IronWustiteBuffer()},
    )


@pytest.fixture(scope="module")
def cho_dicts(cho_output: Output) -> dict[str, dict[str, Any]]:
    return _to_dicts(cho_output)


@pytest.fixture(scope="module")
def h_only_dicts() -> dict[str, dict[str, Any]]:
    """H2 and H gas, so there are no heavy elements (two species avoid the degenerate
    single-species case)"""
    return _to_dicts(_solve((H2_g, H_g), {"H": earth.oceans_to_hydrogen_mass(1)}))


@pytest.fixture(scope="module")
def no_h_output() -> Output:
    """C-O gas without H, so abundances relative to H are not defined"""
    return _solve(
        (O2_g, CO_g, CO2_g),
        {"C": earth.oceans_to_hydrogen_mass(1)},
        activity_constraints={"O2_g": IronWustiteBuffer()},
    )


@pytest.fixture(scope="module")
def no_h_dicts(no_h_output: Output) -> dict[str, dict[str, Any]]:
    return _to_dicts(no_h_output)


def _gas_phase(out: dict[str, Any]) -> dict[str, float]:
    """Gas phase-level outputs of the first batch element"""
    return {key: float(np.ravel(value)[0]) for key, value in out["gas"]["phase"].items()}


def _gas_abundance(dicts: dict[str, dict[str, Any]], key: str) -> dict[str, Optional[dict]]:
    """Gas-phase abundance relative to H (``log_epsilon`` or ``X_H``) of the first batch element
    for each output format, or ``None`` if absent"""
    elements: dict[str, Any] = dicts["natural"]["gas"]["elements"]
    named: dict[str, Any] = dicts["named_arrays"]["gas"]["elements"]
    by_element: dict[str, Any] = {
        name: dicts["elements_species"][name]["gas"] for name in elements["names"]
    }

    return {
        "natural": (
            dict(zip(elements["names"], np.asarray(elements[key])[0])) if key in elements else None
        ),
        "named_arrays": (
            {name: float(np.ravel(v)[0]) for name, v in named[key].items()}
            if key in named
            else None
        ),
        "elements_species": (
            {name: float(np.ravel(v[key])[0]) for name, v in by_element.items()}
            if all(key in v for v in by_element.values())
            else None
        ),
    }


def test_to_dict(cho_output: Output, cho_dicts: dict[str, dict[str, Any]]) -> None:
    """Output formats, the default format, batch expansion, and the solution"""
    for out in cho_dicts.values():
        assert {"gas", "solution", "initial_guess", "solver", "state"} <= set(out)

    natural: dict[str, Any] = cho_dicts["natural"]["gas"]["species"]
    assert "H2O_g" in natural["names"]
    assert natural["partial_pressure"].shape[-1] == len(natural["names"])
    assert "H2O_g" in cho_dicts["named_arrays"]["gas"]["species"]["partial_pressure"]
    assert {"H2O_g", "H", "O"} <= set(cho_dicts["elements_species"])
    assert set(cho_output.to_dict(to_numpy=True)) == set(cho_dicts["named_arrays"])

    expanded: dict[str, Any] = cho_output.to_dict(
        output_format="elements_species", to_numpy=True, expand_to_batch=True, ravel=True
    )
    assert np.asarray(expanded["state"]["temperature"]).shape == (2,)
    assert np.asarray(expanded["gas"]["phase"]["mass"]).shape == (2,)

    assert cho_output.solution.shape == (2, 2 * len(cho_output.parameters.species.species_names))

    # Beartype rejects an invalid Literal before the method body runs, so call the undecorated
    # function to exercise the method's own ValueError
    to_dict = getattr(Output.to_dict, "__wrapped__", Output.to_dict)
    with pytest.raises(ValueError):
        to_dict(cho_output, output_format="not_a_real_format")  # pyright: ignore


def test_compare(cho_output: Output, cho_dicts: dict[str, dict[str, Any]]) -> None:
    """Output.compare matches the output, rejects a mismatch, and compares logarithms"""
    pressure: np.ndarray = cho_dicts["named_arrays"]["gas"]["species"]["partial_pressure"]["H2O_g"]
    pressure = np.ravel(pressure)  # As compare, which expands to the batch and ravels

    def target(value: Any) -> dict[str, Any]:
        return {"gas": {"species": {"partial_pressure": {"H2O_g": value}}}}

    assert cho_output.compare(target(pressure), rtol=1.0e-6, atol=1.0e-6)
    assert not cho_output.compare(target(-1.0), rtol=1.0e-6, atol=1.0e-6)
    assert cho_output.compare(target(pressure), rtol=1.0e-6, atol=1.0e-6, log=True)


def test_quick_look(cho_output: Output, caplog: pytest.LogCaptureFixture) -> None:
    """quick_look logs the output"""
    with caplog.at_level(logging.INFO, logger="atmodeller"):
        cho_output.quick_look()

    assert "Quick look output" in caplog.text


def test_to_dataframes(cho_output: Output) -> None:
    """Dataframes in both formats, and dropping unsuccessful solves"""
    dataframes: dict[str, pd.DataFrame] = cho_output.to_dataframes()
    assert all(isinstance(df, pd.DataFrame) for df in dataframes.values())
    assert any(col.startswith("species.partial_pressure") for col in dataframes["gas"].columns)

    assert isinstance(
        cho_output.to_dataframes(output_format="elements_species")["H2O_g"], pd.DataFrame
    )

    n_success: int = int(np.sum(np.asarray(cho_output.multi_attempt_solution.success)))
    for df in cho_output.to_dataframes(drop_unsuccessful_solves=True).values():
        assert len(df) == n_success


def test_unsuccessful_solves(cho_output: Output, tmp_path: Path) -> None:
    """A model that failed the convergence post-check is dropped and highlighted in Excel, even
    though the solver status reports success"""
    # Emulate a failed post-check of the second model, as reported by the batch retry solver
    output: Output = eqx.tree_at(
        lambda out: out.multi_attempt_solution._attempts, cho_output, jnp.array([1, 0])
    )
    assert np.all(np.asarray(output.multi_attempt_solution.solver_success))

    for df in output.to_dataframes(drop_unsuccessful_solves=True).values():
        assert df.index.tolist() == [0]

    # Excel rows are 1-indexed and row 1 is the header, so the second model is row 3
    for drop, expected_rows in ((False, [3]), (True, [])):
        output.to_excel(
            file_prefix=tmp_path / "out", drop_unsuccessful_solves=drop, append_version=False
        )
        for sheet in load_workbook(tmp_path / "out.xlsx").worksheets:
            highlighted_rows: list[int] = [
                number
                for number, row in enumerate(sheet.iter_rows(), start=1)
                if row[0].fill.fgColor.rgb == "00FFFF00"
            ]
            assert highlighted_rows == expected_rows


def _read_excel(path: Path) -> dict[str, pd.DataFrame]:
    return pd.read_excel(path, sheet_name=None, index_col=0)


def _read_pickle(path: Path) -> Any:
    with open(path, "rb") as handle:
        return pickle.load(handle)


@pytest.mark.parametrize(
    "export, extension, read_back",
    [
        pytest.param(Output.to_excel, "xlsx", _read_excel, id="excel"),
        pytest.param(Output.to_pickle, "pkl", _read_pickle, id="pickle"),
    ],
)
def test_exports(
    cho_output: Output,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    export,
    extension: str,
    read_back,
) -> None:
    """File names with and without the version, an existing extension, and a relative prefix"""
    export(cho_output, file_prefix=tmp_path / "out")
    assert "gas" in read_back(tmp_path / f"out_v{__version__}.{extension}")

    export(cho_output, file_prefix=tmp_path / "no_version", append_version=False)
    assert (tmp_path / f"no_version.{extension}").exists()

    export(cho_output, file_prefix=tmp_path / f"suffixed.{extension}", append_version=False)
    assert (tmp_path / f"suffixed.{extension}").exists()
    assert not (tmp_path / f"suffixed.{extension}.{extension}").exists()

    monkeypatch.chdir(tmp_path)
    export(cho_output, file_prefix="relative", append_version=False)
    assert (tmp_path / f"relative.{extension}").exists()


def test_metallicity(
    cho_dicts: dict[str, dict[str, Any]],
    h_only_dicts: dict[str, dict[str, Any]],
    no_h_dicts: dict[str, dict[str, Any]],
) -> None:
    """Gas metallicity in all formats, with heavy elements, without them, and without H"""
    named: dict[str, Any] = cho_dicts["named_arrays"]["gas"]["elements"]
    n: dict[str, float] = {key: float(np.ravel(v)[0]) for key, v in named["number_moles"].items()}
    m: dict[str, float] = {key: float(np.ravel(v)[0]) for key, v in named["mass"].items()}
    metals_to_hydrogen: float = (n["C"] + n["O"]) / n["H"]
    solar_metals_to_hydrogen: float = sum(10 ** (solar.log_epsilon[e] - 12) for e in ("C", "O"))

    for fmt in OUTPUT_FORMATS:
        # C and O are metals
        gas: dict[str, float] = _gas_phase(cho_dicts[fmt])
        assert "metallicity_by_moles" not in gas
        assert gas["metallicity_by_mass"] > 0
        assert np.isclose(gas["metals_to_hydrogen"], metals_to_hydrogen)
        assert np.isclose(gas["metals_to_hydrogen_by_mass"], (m["C"] + m["O"]) / m["H"])
        assert np.isclose(gas["M_H"], np.log10(metals_to_hydrogen / solar_metals_to_hydrogen))

        # Only H, so no metals and [M/H] is undefined
        gas = _gas_phase(h_only_dicts[fmt])
        assert gas["metallicity_by_mass"] == 0.0
        assert gas["metals_to_hydrogen"] == 0.0
        assert gas["metals_to_hydrogen_by_mass"] == 0.0
        assert np.isnan(gas["M_H"])

        # No H, so all elements are metals and quantities relative to H are not defined
        gas = _gas_phase(no_h_dicts[fmt])
        assert np.isclose(gas["metallicity_by_mass"], 1.0)
        assert not {"metals_to_hydrogen", "metals_to_hydrogen_by_mass", "M_H"} & set(gas)


def test_abundances_relative_to_hydrogen(
    cho_dicts: dict[str, dict[str, Any]],
    no_h_output: Output,
    no_h_dicts: dict[str, dict[str, Any]],
) -> None:
    """log epsilon and [X/H] in all formats, and their absence without H"""
    moles: np.ndarray = np.array([[2.0, 0.02, 0.0], [4.0, 4.0, 0.4]])
    result: np.ndarray = np.asarray(log_epsilon(jnp.asarray(moles), 0))
    np.testing.assert_allclose(result[:, 0], 12.0)
    np.testing.assert_allclose([result[0, 1], *result[1, 1:]], [10.0, 12.0, 11.0])
    assert result[0, 2] == -np.inf  # An absent element

    named: dict[str, Any] = cho_dicts["named_arrays"]["gas"]["elements"]["number_moles"]
    n_h, n_c = float(np.ravel(named["H"])[0]), float(np.ravel(named["C"])[0])
    x_h: dict[str, Optional[dict]] = _gas_abundance(cho_dicts, "X_H")
    for fmt, values in _gas_abundance(cho_dicts, "log_epsilon").items():
        relative: Optional[dict] = x_h[fmt]
        assert values is not None and relative is not None
        assert values["H"] == 12.0
        assert np.isclose(values["C"], np.log10(n_c / n_h) + 12)
        for element, value in values.items():
            assert np.isclose(relative[element], value - solar.log_epsilon[element])

    for key in ("log_epsilon", "X_H"):
        assert all(values is None for values in _gas_abundance(no_h_dicts, key).values())
    columns: pd.Index = no_h_output.to_dataframes()["gas"].columns
    assert not any("log_epsilon" in column or "X_H" in column for column in columns)


def test_matches_v1(cho_dicts: dict[str, dict[str, Any]]) -> None:
    """log epsilon and [M/H] match ``logarithmic_abundance`` from v1 (commit 77e81aac) for the same
    gas-only network. H and C are set by the mass constraints, so agree to round-off; O depends on
    the speciation, so agrees to the solver tolerance."""
    target: dict[str, float] = {"H": 12.0, "C": 10.923865345578943, "O": 11.54801489087738}
    for values in _gas_abundance(cho_dicts, "log_epsilon").values():
        assert values is not None
        for element, atol in (("H", 1e-12), ("C", 1e-12), ("O", 1e-7)):
            assert np.isclose(values[element], target[element], rtol=0, atol=atol)

    metals_to_hydrogen: float = 10 ** (target["C"] - 12) + 10 ** (target["O"] - 12)
    solar_metals_to_hydrogen: float = sum(10 ** (solar.log_epsilon[e] - 12) for e in ("C", "O"))
    gas: dict[str, float] = _gas_phase(cho_dicts["named_arrays"])
    assert np.isclose(gas["metals_to_hydrogen"], metals_to_hydrogen, rtol=1e-7)
    assert np.isclose(
        gas["M_H"], np.log10(metals_to_hydrogen / solar_metals_to_hydrogen), rtol=0, atol=1e-7
    )


def test_solar_abundances() -> None:
    """Spot checks against Asplund et al. (2021, Table 2), and NaN for a missing element"""
    for element, value in {"H": 12.0, "C": 8.46, "N": 7.83, "O": 8.69, "S": 7.12}.items():
        assert solar.log_epsilon[element] == value

    values: np.ndarray = solar.get_log_epsilon(("H", "C", "As"))
    assert np.allclose(values[:2], [12.0, 8.46])
    assert np.isnan(values[2])


def test_dictionary_helpers() -> None:
    """flatten_dictionary and recursively_merge_dictionaries"""
    assert flatten_dictionary({"a": {"b": {"c": 1}, "d": 2}, "e": 3}) == {
        "a.b.c": 1,
        "a.d": 2,
        "e": 3,
    }
    assert flatten_dictionary({}) == {}
    assert flatten_dictionary({"b": 1}, parent_key="a") == {"a.b": 1}

    d1: dict[str, Any] = {"a": {"x": 1, "y": 2}, "b": {"x": 1}, "c": 1}
    d2: dict[str, Any] = {"a": {"y": 3, "z": 4}, "b": 5, "d": 2}
    assert recursively_merge_dictionaries(d1, d2) == {
        "a": {"x": 1, "y": 3, "z": 4},
        "b": 5,
        "c": 1,
        "d": 2,
    }
    # The inputs are not mutated
    assert d1 == {"a": {"x": 1, "y": 2}, "b": {"x": 1}, "c": 1}
    assert d2 == {"a": {"y": 3, "z": 4}, "b": 5, "d": 2}
